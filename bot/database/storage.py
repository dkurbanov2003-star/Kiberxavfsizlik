"""
storage.py — Ma'lumotlarni saqlash qatlami.

Ikki joyga yozamiz:
  1) SQLite (bot.db) — operatsion ma'lumot: kim qatnashgani, ishtirokchi #ID,
     jins, yosh, ball. Bu ID raqamining takrorlanmasligini kafolatlaydi.
  2) Excel (malumotlar.xlsx) — SPSS tahliliga TAYYOR jadval:
     ID | Jinsi | Yoshi | Ball | Sana

Nega ikkalasi? SQLite ma'lumotni ishonchli boshqaradi (ID hisoblagichi),
Excel esa to'g'ridan-to'g'ri SPSS'ga import qilinadi.

Diqqat: aiogram bir jarayonda (async) ishlagani uchun, fayl buzilmasligi
uchun barcha yozishni asyncio.Lock bilan himoyalaymiz.
"""

import asyncio
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from bot.config import BASE_DIR

DB_PATH = BASE_DIR / "bot.db"

# Bir vaqtda faqat bitta yozuv amali bajarilishini kafolatlaydigan qulf
_lock = asyncio.Lock()
_conn: sqlite3.Connection | None = None


def _connection() -> sqlite3.Connection:
    """SQLite ulanishini bir marta ochib, qayta ishlatamiz."""
    global _conn
    if _conn is None:
        _conn = sqlite3.connect(DB_PATH, check_same_thread=False)
        _conn.row_factory = sqlite3.Row
    return _conn


async def init_db() -> None:
    """Bot ishga tushganda jadvalni yaratadi (agar yo'q bo'lsa)."""
    async with _lock:
        conn = _connection()
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS participants (
                telegram_id     INTEGER PRIMARY KEY,
                participant_id  INTEGER UNIQUE,
                gender          TEXT,
                age             TEXT,
                score           INTEGER,
                created_at      TEXT
            )
            """
        )
        conn.commit()


async def register_participant(telegram_id: int, gender: str, age: str) -> int:
    """
    Foydalanuvchini ro'yxatga oladi va unga takrorlanmas #ID beradi.
    Agar u avval qatnashgan bo'lsa, o'sha ID qaytariladi (dublikat bo'lmaydi).
    """
    async with _lock:
        conn = _connection()
        row = conn.execute(
            "SELECT participant_id FROM participants WHERE telegram_id = ?",
            (telegram_id,),
        ).fetchone()

        if row is not None:
            # Avval qatnashgan — jins/yoshni yangilab, eski ID'ni qaytaramiz
            conn.execute(
                "UPDATE participants SET gender = ?, age = ? WHERE telegram_id = ?",
                (gender, age, telegram_id),
            )
            conn.commit()
            return row["participant_id"]

        # Yangi ishtirokchi — keyingi tartib raqamini beramiz
        max_row = conn.execute(
            "SELECT COALESCE(MAX(participant_id), 0) AS m FROM participants"
        ).fetchone()
        new_id = max_row["m"] + 1

        conn.execute(
            """
            INSERT INTO participants (telegram_id, participant_id, gender, age, created_at)
            VALUES (?, ?, ?, ?, ?)
            """,
            (telegram_id, new_id, gender, age, datetime.now().isoformat(timespec="seconds")),
        )
        conn.commit()
        return new_id


async def save_result(telegram_id: int, score: int, excel_path: Path) -> tuple[int, str, str]:
    """
    Test ballini saqlaydi va SPSS uchun Excel qatorini yozadi.
    Qaytaradi: (participant_id, gender, age).
    """
    async with _lock:
        conn = _connection()
        conn.execute(
            "UPDATE participants SET score = ? WHERE telegram_id = ?",
            (score, telegram_id),
        )
        conn.commit()

        row = conn.execute(
            "SELECT participant_id, gender, age FROM participants WHERE telegram_id = ?",
            (telegram_id,),
        ).fetchone()

        participant_id = row["participant_id"]
        gender = row["gender"] or "—"
        age = row["age"] or "—"

        _write_to_excel(excel_path, participant_id, gender, age, score)
        return participant_id, gender, age


def _write_to_excel(path: Path, participant_id: int, gender: str, age: str, score: int) -> None:
    """Excel faylga bitta qator yozadi (mavjud ID'ni yangilaydi yoki yangi qo'shadi)."""
    id_str = f"#{participant_id:05d}"
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")

    if Path(path).exists():
        workbook = load_workbook(path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Natijalar"
        sheet.append(["ID", "Jinsi", "Yoshi", "Ball", "Sana"])  # sarlavha

    # Shu ID allaqachon bormi? Bo'lsa — yangilaymiz (qayta test topshirsa)
    updated = False
    for excel_row in sheet.iter_rows(min_row=2):
        if excel_row[0].value == id_str:
            excel_row[1].value = gender
            excel_row[2].value = age
            excel_row[3].value = score
            excel_row[4].value = now_str
            updated = True
            break

    if not updated:
        sheet.append([id_str, gender, age, score, now_str])

    workbook.save(path)
